import csv
from argparse import ArgumentParser
from decimal import Decimal
from pathlib import Path
from typing import Any

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from whimo.transactions.qr_cards import DEFAULT_BOUNDING_BOX, FarmRow, build_cards, render_cards_pdf


class Command(BaseCommand):
    help = "Generate printable farm QR cards from a cooperative spreadsheet"

    def add_arguments(self, parser: ArgumentParser) -> None:
        parser.add_argument("input", type=Path, help="Spreadsheet with the farmer ids and their coordinates")
        parser.add_argument("output", type=Path, help="PDF file to write the cards to")
        parser.add_argument("--rejects", type=Path, help="CSV file to write the rejected rows to")
        parser.add_argument("--sheet", help="Sheet to read, the first one by default")
        parser.add_argument("--id-column", default="farmer id")
        parser.add_argument("--latitude-column", default="latitude")
        parser.add_argument("--longitude-column", default="longitude")
        parser.add_argument("--bounding-box", help="min_longitude,min_latitude,max_longitude,max_latitude")
        parser.add_argument("--columns", type=int, default=3)
        parser.add_argument("--rows", type=int, default=4)
        parser.add_argument("--error-level", default="M", choices=["L", "M", "Q", "H"])

    def handle(self, *_: Any, **options: Any) -> None:
        rows = self._read_rows(options)
        bounding_box = self._bounding_box(options["bounding_box"])
        cards, rejected = build_cards(rows, bounding_box=bounding_box)

        pages = render_cards_pdf(
            cards,
            options["output"],
            columns=options["columns"],
            rows=options["rows"],
            error_level=options["error_level"],
        )

        if rejected and (rejects_path := options["rejects"]):
            with rejects_path.open("w", newline="") as rejects_file:
                writer = csv.writer(rejects_file)
                writer.writerow(["Row", "Farmer ID", "Reason"])
                writer.writerows([(item.number, item.farmer_id, item.reason) for item in rejected])

        self.stdout.write(f"{len(cards)} cards on {pages} pages, {len(rejected)} rows rejected")
        if rejected and not options["rejects"]:
            for item in rejected:
                self.stdout.write(f"row {item.number}: {item.farmer_id or '-'}: {item.reason}")

    def _read_rows(self, options: dict[str, Any]) -> list[FarmRow]:
        if not options["input"].exists():
            raise CommandError(f"{options['input']} does not exist")

        workbook = load_workbook(options["input"], read_only=True, data_only=True)
        sheet = workbook[options["sheet"]] if options["sheet"] else workbook.worksheets[0]

        sheet_rows = sheet.iter_rows(values_only=True)
        header = [str(cell).strip().lower() if cell is not None else "" for cell in next(sheet_rows, ())]
        indexes = [
            self._column_index(header, options[name]) for name in ("id_column", "latitude_column", "longitude_column")
        ]

        rows = [
            FarmRow(
                number=number,
                farmer_id=self._cell(row, indexes[0]),
                latitude=self._cell(row, indexes[1]),
                longitude=self._cell(row, indexes[2]),
            )
            for number, row in enumerate(sheet_rows, start=2)
            if any(cell is not None for cell in row)
        ]
        workbook.close()

        return rows

    @staticmethod
    def _column_index(header: list[str], name: str) -> int:
        try:
            return header.index(name.strip().lower())
        except ValueError as err:
            raise CommandError(f"column {name!r} is missing, the sheet has {header}") from err

    @staticmethod
    def _cell(row: tuple[Any, ...], index: int) -> Any:
        return row[index] if index < len(row) else None

    @staticmethod
    def _bounding_box(value: str | None) -> tuple[Decimal, Decimal, Decimal, Decimal]:
        if not value:
            return DEFAULT_BOUNDING_BOX

        try:
            min_longitude, min_latitude, max_longitude, max_latitude = (Decimal(part) for part in value.split(","))
        except (ValueError, ArithmeticError) as err:
            raise CommandError("bounding box must be min_longitude,min_latitude,max_longitude,max_latitude") from err

        return min_longitude, min_latitude, max_longitude, max_latitude
